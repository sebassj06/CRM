import csv
import io
from datetime import date, datetime
import openpyxl
from core.clientes import agregar_cliente, obtener_cliente_por_email, email_valido


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%Y-%m-%d")
    return str(valor).strip()


def _filas_desde_csv(archivo):
    contenido = archivo.stream.read().decode("utf-8-sig")
    lector = csv.DictReader(io.StringIO(contenido))
    return list(lector)


def _filas_desde_excel(archivo):
    libro = openpyxl.load_workbook(archivo.stream, data_only=True)
    hoja = libro.active
    filas_crudas = list(hoja.iter_rows(values_only=True))

    if not filas_crudas:
        return []

    encabezados = [_texto(valor) for valor in filas_crudas[0]]

    filas = []
    for fila_valores in filas_crudas[1:]:
        fila_dict = {}
        for indice, encabezado in enumerate(encabezados):
            fila_dict[encabezado] = fila_valores[indice] if indice < len(fila_valores) else None
        filas.append(fila_dict)
    return filas


def importar_clientes_desde_archivo(archivo, agencia_id):
    nombre_archivo = archivo.filename.lower()

    if nombre_archivo.endswith(".xlsx"):
        filas = _filas_desde_excel(archivo)
    elif nombre_archivo.endswith(".csv"):
        filas = _filas_desde_csv(archivo)
    else:
        return None

    importados = 0
    duplicados = 0
    invalidos = 0

    for fila in filas:
        nombre = _texto(fila.get("nombre"))
        email = _texto(fila.get("email"))
        telefono = _texto(fila.get("telefono"))
        empresa = _texto(fila.get("empresa"))
        notas = _texto(fila.get("notas"))
        cliente_desde = _texto(fila.get("cliente_desde"))

        if not nombre or not email or not email_valido(email):
            invalidos += 1
            continue

        if obtener_cliente_por_email(email, agencia_id):
            duplicados += 1
            continue

        agregar_cliente(nombre, email, telefono, empresa, notas, agencia_id, cliente_desde or None)
        importados += 1

    return {"importados": importados, "duplicados": duplicados, "invalidos": invalidos}


def generar_plantilla_clientes():
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Clientes"
    hoja.append(["nombre", "email", "telefono", "empresa", "notas", "cliente_desde"])
    hoja.append(["Ejemplo Cliente", "ejemplo@cliente.com", "1122334455", "Ejemplo SA", "Cliente de ejemplo, podés borrar esta fila", "2024-01-15"])

    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer